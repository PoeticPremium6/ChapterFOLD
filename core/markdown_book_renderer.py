"""Render edited ChapterFOLD Markdown back into printable outputs.

This module is intentionally conservative. It gives ChapterFOLD a stable
"pre-edit" workflow:

    EPUB -> Editable.md -> user edits Markdown -> PDF/DOCX

It does not try to be a full Markdown implementation. It supports the subset
ChapterFOLD currently emits: titles, author lines, headings, paragraphs,
horizontal rules, simple emphasis, and fenced/preformatted blocks.
"""
from __future__ import annotations

from core.markdown_sanitizer import sanitize_chapterfold_markdown_for_render

from dataclasses import dataclass
from html import escape
from pathlib import Path
import re
from core.toc_rendering import apply_toc_mode_to_markdown
from core.font_policy import UNICODE_PDF_FONT_STACK


@dataclass(frozen=True)
class MarkdownBookMetadata:
    title: str
    author: str = ""


@dataclass(frozen=True)
class MarkdownRenderResult:
    success: bool
    output_files: list[str]
    warnings: list[str]
    error: str | None
    title: str
    author: str
    markdown_chars: int
    estimated_word_count: int


_TITLE_RE = re.compile(r"^\s*#\s+(.+?)\s*$")
_AUTHOR_RE = re.compile(r"^\s*_By\s+(.+?)_\s*$", re.IGNORECASE)
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
_HR_RE = re.compile(r"^\s{0,3}([-*_])(?:\s*\1){2,}\s*$")


def slugify_filename(value: str, fallback: str = "book") -> str:
    """Return a filesystem-safe stem."""
    value = re.sub(r"[\\/:*?\"<>|]+", " ", value)
    value = re.sub(r"\s+", " ", value).strip(" ._-")
    return value or fallback


def parse_markdown_metadata(markdown_text: str, *, fallback_title: str = "Edited Book") -> MarkdownBookMetadata:
    """Infer title/author from ChapterFOLD-style Markdown."""
    title = fallback_title
    author = ""

    lines = markdown_text.splitlines()
    for line in lines[:40]:
        m = _TITLE_RE.match(line)
        if m:
            title = m.group(1).strip()
            break

    for line in lines[:60]:
        m = _AUTHOR_RE.match(line)
        if m:
            author = m.group(1).strip()
            break

    return MarkdownBookMetadata(title=title, author=author)


def _inline_markdown_to_html(text: str) -> str:
    """Very small inline Markdown renderer for ChapterFOLD output."""
    text = escape(text)
    text = re.sub(r"\\\*", "*", text)
    text = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", text)
    text = re.sub(r"(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)", r"<em>\1</em>", text)
    text = re.sub(r"_(?!\s)(.+?)(?<!\s)_", r"<em>\1</em>", text)
    text = text.replace("  \n", "<br>\n")
    return text


def markdown_to_html_body(markdown_text: str) -> str:
    """Convert a practical subset of Markdown to an HTML body."""
    lines = markdown_text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    out: list[str] = []
    paragraph: list[str] = []
    in_fence = False
    fence_lines: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph
        if paragraph:
            joined = " ".join(part.strip() for part in paragraph if part.strip())
            if joined:
                out.append(f"<p>{_inline_markdown_to_html(joined)}</p>")
            paragraph = []

    def flush_fence() -> None:
        nonlocal fence_lines
        out.append("<pre>" + escape("\n".join(fence_lines)).strip("\n") + "</pre>")
        fence_lines = []

    for raw in lines:
        line = raw.rstrip("\n")

        if line.strip().startswith("```"):
            flush_paragraph()
            if in_fence:
                flush_fence()
                in_fence = False
            else:
                in_fence = True
                fence_lines = []
            continue

        if in_fence:
            fence_lines.append(line)
            continue

        if not line.strip():
            flush_paragraph()
            continue

        if _HR_RE.match(line):
            flush_paragraph()
            out.append("<hr>")
            continue

        heading = _HEADING_RE.match(line)
        if heading:
            flush_paragraph()
            level = min(len(heading.group(1)), 6)
            text = _inline_markdown_to_html(heading.group(2).strip())
            out.append(f"<h{level}>{text}</h{level}>")
            continue

        symbol_count = sum(1 for ch in line if ch in "*=+_|<>[]/\\#~")
        if len(line.strip()) >= 8 and symbol_count / max(len(line.strip()), 1) > 0.55:
            flush_paragraph()
            out.append(f"<pre>{escape(line.strip())}</pre>")
            continue

        paragraph.append(line)

    flush_paragraph()
    if in_fence:
        flush_fence()

    return "\n".join(out)


def build_print_html(markdown_text: str, *, title: str | None = None, author: str | None = None) -> str:
    metadata = parse_markdown_metadata(markdown_text)
    title = title or metadata.title
    author = author if author is not None else metadata.author
    body = markdown_to_html_body(markdown_text)

    title_html = escape(title)
    font_stack = UNICODE_PDF_FONT_STACK
    author_html = escape(author) if author else ""

    return f"""<!doctype html>
<html>
<head>
<meta charset="utf-8">
<title>{title_html}</title>
<style>
@page {{
  size: 5.5in 8.5in;
  margin: 0.7in 0.62in 0.78in 0.62in;
  @bottom-center {{ content: counter(page); font-size: 9pt; color: #555; }}
}}
body {{
  font-family: {font_stack};
  font-size: 11pt;
  line-height: 1.32;
  color: #111;
}}
.title-page {{
  break-after: page;
  text-align: center;
  padding-top: 30%;
}}
.title-page h1 {{
  font-size: 24pt;
  line-height: 1.15;
  margin: 0 0 1.5em 0;
}}
.title-page .author {{
  font-size: 14pt;
}}
h1, h2, h3, h4 {{
  break-after: avoid;
  line-height: 1.2;
}}
h1 {{
  font-size: 20pt;
  text-align: center;
  margin-top: 2.5em;
}}
h2 {{
  font-size: 15pt;
  text-align: center;
  margin-top: 2.2em;
}}
h3 {{
  font-size: 13pt;
  margin-top: 1.6em;
}}
p {{
  text-indent: 1.15em;
  margin: 0 0 0.25em 0;
  widows: 2;
  orphans: 2;
}}
h1 + p, h2 + p, h3 + p, hr + p, pre + p {{
  text-indent: 0;
}}
hr {{
  border: 0;
  text-align: center;
  margin: 1.2em 0;
}}
hr::after {{
  content: "* * *";
  letter-spacing: 0.2em;
}}
pre {{
  font-family: "DejaVu Sans Mono", monospace;
  font-size: 8.5pt;
  line-height: 1.15;
  white-space: pre-wrap;
  margin: 0.8em 0;
  text-indent: 0;
}}
</style>
</head>
<body>
<section class="title-page">
  <h1>{title_html}</h1>
  {f'<div class="author">{author_html}</div>' if author_html else ''}
</section>
{body}
</body>
</html>
"""


def estimate_word_count(markdown_text: str) -> int:
    return len(re.findall(r"\b[\w’'-]+\b", markdown_text))


def render_markdown_book(
    input_markdown: str | Path,
    output_dir: str | Path,
    *,
    title: str | None = None,
    author: str | None = None,
    export_pdf: bool = True,
    export_docx: bool = False,
    copy_markdown: bool = True,
    toc_mode: str = "keep",
) -> MarkdownRenderResult:
    """Render edited Markdown into book outputs."""
    input_path = Path(input_markdown)
    output_path = Path(output_dir)
    output_files: list[str] = []
    warnings: list[str] = []

    try:
        if not input_path.exists():
            raise FileNotFoundError(f"Markdown file does not exist: {input_path}")
        if input_path.suffix.lower() not in {".md", ".markdown", ".txt"}:
            warnings.append("Input does not use a .md/.markdown/.txt extension; rendering anyway.")

        markdown_text = input_path.read_text(encoding="utf-8")
        toc_result = apply_toc_mode_to_markdown(markdown_text, toc_mode)
        if hasattr(toc_result, "markdown"):
            markdown_text = toc_result.markdown
        elif isinstance(toc_result, tuple):
            markdown_text = toc_result[0]
        else:
            markdown_text = str(toc_result)
        metadata = parse_markdown_metadata(markdown_text, fallback_title=input_path.stem)
        final_title = title or metadata.title or input_path.stem
        final_author = author if author is not None else metadata.author
        output_path.mkdir(parents=True, exist_ok=True)

        stem_parts = []
        if final_author:
            stem_parts.append(slugify_filename(final_author))
        stem_parts.append(slugify_filename(final_title))
        stem_parts.append("Rendered from Markdown")
        stem = " - ".join(stem_parts)

        if copy_markdown:
            md_out = output_path / f"{stem} - Editable.md"
            md_out.write_text(markdown_text, encoding="utf-8")
            output_files.append(str(md_out))

        if export_pdf:
            html = build_print_html(markdown_text, title=final_title, author=final_author)
            pdf_out = output_path / f"{stem} - Interior.pdf"
            try:
                from weasyprint import HTML
            except Exception as exc:  # pragma: no cover
                raise RuntimeError(f"WeasyPrint is required for PDF rendering: {exc}") from exc
            HTML(string=html, base_url=str(input_path.parent)).write_pdf(str(pdf_out))
            output_files.append(str(pdf_out))

        if export_docx:
            docx_out = output_path / f"{stem} - Editable.docx"
            try:
                from docx import Document
            except Exception as exc:  # pragma: no cover
                raise RuntimeError(f"python-docx is required for DOCX export: {exc}") from exc

            document = Document()
            document.add_heading(final_title, level=0)
            if final_author:
                document.add_paragraph(f"By {final_author}")

            for line in markdown_text.splitlines():
                stripped = line.strip()
                if not stripped:
                    continue
                if stripped.startswith("#"):
                    level = min(len(stripped) - len(stripped.lstrip("#")), 4)
                    document.add_heading(stripped.lstrip("#").strip(), level=max(level, 1))
                elif _HR_RE.match(stripped):
                    document.add_paragraph("* * *")
                elif stripped.startswith("_By ") and stripped.endswith("_"):
                    continue
                else:
                    document.add_paragraph(stripped)
            document.save(str(docx_out))
            output_files.append(str(docx_out))

        return MarkdownRenderResult(
            success=True,
            output_files=output_files,
            warnings=warnings,
            error=None,
            title=final_title,
            author=final_author,
            markdown_chars=len(markdown_text),
            estimated_word_count=estimate_word_count(markdown_text),
        )

    except Exception as exc:
        return MarkdownRenderResult(
            success=False,
            output_files=output_files,
            warnings=warnings,
            error=str(exc),
            title=title or "",
            author=author or "",
            markdown_chars=0,
            estimated_word_count=0,
        )

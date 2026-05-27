from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from pypdf import PdfReader
from docx import Document


@dataclass
class PdfExtractedPage:
    page_number: int
    text: str
    char_count: int


@dataclass
class PdfTextQualityReport:
    page_count: int
    total_chars: int
    average_chars_per_page: float
    empty_pages: list[int]
    near_empty_pages: list[int]
    likely_scanned_pages: list[int]
    garbled_pages: list[int]
    repeated_line_candidates: list[str]
    ocr_needed: str
    confidence: str


@dataclass
class PdfTextExtractionResult:
    success: bool
    input_pdf: Path
    page_count: int
    pages: list[PdfExtractedPage] = field(default_factory=list)
    total_chars: int = 0
    quality_report: PdfTextQualityReport | None = None
    markdown_path: Path | None = None
    docx_path: Path | None = None
    warnings: list[str] = field(default_factory=list)
    error: str | None = None



def _looks_garbled(value: str) -> bool:
    if not value.strip():
        return False

    chars = [c for c in value if not c.isspace()]
    if not chars:
        return False

    replacement_count = value.count("\ufffd")
    control_count = sum(1 for c in chars if ord(c) < 32)
    alpha_count = sum(1 for c in chars if c.isalpha())
    punctuation_count = sum(1 for c in chars if c in "{}[]<>|~^_=+*#")

    if replacement_count >= 3:
        return True
    if control_count / max(len(chars), 1) > 0.05:
        return True
    if len(chars) > 80 and alpha_count / len(chars) < 0.25:
        return True
    if len(chars) > 80 and punctuation_count / len(chars) > 0.20:
        return True

    return False


def _repeated_line_candidates(pages: list[PdfExtractedPage], *, min_repeats: int = 3) -> list[str]:
    counts: dict[str, int] = {}

    for page in pages:
        for line in page.text.splitlines():
            cleaned = " ".join(line.strip().split())
            if 4 <= len(cleaned) <= 120:
                counts[cleaned] = counts.get(cleaned, 0) + 1

    repeated = [
        line
        for line, count in counts.items()
        if count >= min_repeats
    ]

    repeated.sort(key=lambda line: (-counts[line], line.lower()))
    return repeated[:20]


def build_pdf_text_quality_report(result: PdfTextExtractionResult) -> PdfTextQualityReport:
    page_count = result.page_count
    total_chars = result.total_chars
    average = total_chars / page_count if page_count else 0.0

    empty_pages = [page.page_number for page in result.pages if page.char_count == 0]
    near_empty_pages = [page.page_number for page in result.pages if 0 < page.char_count < 40]
    likely_scanned_pages = [page.page_number for page in result.pages if page.char_count < 10]
    garbled_pages = [page.page_number for page in result.pages if _looks_garbled(page.text)]
    repeated = _repeated_line_candidates(result.pages)

    scanned_ratio = len(likely_scanned_pages) / page_count if page_count else 0.0
    garbled_ratio = len(garbled_pages) / page_count if page_count else 0.0

    if total_chars == 0 or scanned_ratio > 0.60:
        ocr_needed = "yes"
        confidence = "poor"
    elif scanned_ratio > 0.20 or garbled_ratio > 0.15 or average < 250:
        ocr_needed = "maybe"
        confidence = "mixed"
    else:
        ocr_needed = "no"
        confidence = "good"

    return PdfTextQualityReport(
        page_count=page_count,
        total_chars=total_chars,
        average_chars_per_page=average,
        empty_pages=empty_pages,
        near_empty_pages=near_empty_pages,
        likely_scanned_pages=likely_scanned_pages,
        garbled_pages=garbled_pages,
        repeated_line_candidates=repeated,
        ocr_needed=ocr_needed,
        confidence=confidence,
    )


def extract_pdf_text_pages(input_pdf: str | Path) -> PdfTextExtractionResult:
    input_pdf = Path(input_pdf)

    if input_pdf.suffix.lower() != ".pdf":
        raise ValueError("Input file must be a .pdf file.")
    if not input_pdf.exists():
        raise FileNotFoundError(f"Input PDF does not exist: {input_pdf}")

    reader = PdfReader(str(input_pdf))
    extracted_pages: list[PdfExtractedPage] = []
    warnings: list[str] = []

    for index, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            text = ""
            warnings.append(f"Page {index}: text extraction failed: {exc}")

        normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
        extracted_pages.append(
            PdfExtractedPage(
                page_number=index,
                text=normalized,
                char_count=len(normalized),
            )
        )

    total_chars = sum(page.char_count for page in extracted_pages)

    if total_chars == 0:
        warnings.append("No extractable text found. This PDF may be scanned or image-only.")

    result = PdfTextExtractionResult(
        success=True,
        input_pdf=input_pdf,
        page_count=len(reader.pages),
        pages=extracted_pages,
        total_chars=total_chars,
        warnings=warnings,
    )
    result.quality_report = build_pdf_text_quality_report(result)

    if result.quality_report.ocr_needed == "yes":
        warnings.append("OCR is likely needed for this PDF.")
    elif result.quality_report.ocr_needed == "maybe":
        warnings.append("OCR may be needed for parts of this PDF.")

    return result


def safe_pdf_text_stem(path: Path) -> str:
    stem = path.stem.strip() or "pdf-text"
    for char in '<>:"/\\|?*':
        stem = stem.replace(char, "-")
    return " ".join(stem.split())



def _line_looks_like_heading(line: str) -> bool:
    value = " ".join((line or "").strip().split())
    if not value:
        return False

    if len(value) > 90:
        return False

    lower = value.lower()
    if lower.startswith(("http://", "https://", "www.")):
        return False

    # Conservative heading guesses only.
    if value.startswith(("Chapter ", "CHAPTER ", "Part ", "PART ", "Appendix ", "APPENDIX ")):
        return True

    if len(value) <= 70 and value.isupper() and any(char.isalpha() for char in value):
        return True

    return False


def _is_probable_page_number_line(line: str) -> bool:
    value = (line or "").strip()
    return value.isdigit() and 1 <= len(value) <= 4


def _join_wrapped_lines(lines: list[str]) -> list[str]:
    paragraphs: list[str] = []
    buffer: list[str] = []

    def flush() -> None:
        nonlocal buffer
        if buffer:
            paragraphs.append(" ".join(part.strip() for part in buffer if part.strip()))
            buffer = []

    for raw_line in lines:
        line = raw_line.strip()

        if not line:
            flush()
            continue

        if _is_probable_page_number_line(line):
            continue

        if _line_looks_like_heading(line):
            flush()
            paragraphs.append(f"## {line}")
            continue

        if buffer and buffer[-1].endswith("-") and line and line[0].islower():
            buffer[-1] = buffer[-1][:-1] + line
            continue

        buffer.append(line)

        if line.endswith((".", "?", "!", ":", ";", ".”", "’", ")")):
            flush()

    flush()
    return paragraphs


def page_text_to_editable_markdown(page: PdfExtractedPage) -> str:
    if not page.text.strip():
        return "[No extractable text found on this page.]"

    cleaned_lines = [line.rstrip() for line in page.text.splitlines()]
    paragraphs = _join_wrapped_lines(cleaned_lines)

    if not paragraphs:
        return page.text.strip()

    return "\n\n".join(paragraphs)


def pdf_text_to_markdown(result: PdfTextExtractionResult) -> str:
    lines: list[str] = []

    lines.append(f"# Editable draft extracted from: {result.input_pdf.stem}")
    lines.append("")
    lines.append("> Experimental PDF text recovery. This Markdown is an editable draft, not a polished book conversion.")
    lines.append("> OCR, layout-perfect reconstruction, footnotes, tables, images, columns, and advanced reflow are not included in this mode.")
    lines.append("")
    lines.append(f"- Source PDF: `{result.input_pdf}`")
    lines.append(f"- Pages: {result.page_count}")
    lines.append(f"- Extracted characters: {result.total_chars}")

    if result.quality_report:
        report = result.quality_report
        lines.append(f"- Extraction confidence: {report.confidence}")
        lines.append(f"- OCR needed: {report.ocr_needed}")
        lines.append("")

        lines.append("## Extraction Quality Report")
        lines.append("")
        lines.append(f"- Average characters per page: {report.average_chars_per_page:.1f}")
        lines.append(f"- Empty pages: {len(report.empty_pages)}")
        lines.append(f"- Near-empty pages: {len(report.near_empty_pages)}")
        lines.append(f"- Likely scanned/image-only pages: {len(report.likely_scanned_pages)}")
        lines.append(f"- Garbled-text pages: {len(report.garbled_pages)}")

        if report.repeated_line_candidates:
            lines.append("- Repeated line candidates:")
            for candidate in report.repeated_line_candidates[:10]:
                lines.append(f"  - {candidate}")
        lines.append("")
    else:
        lines.append("")

    if result.warnings:
        lines.append("## Warnings")
        lines.append("")
        for warning in result.warnings:
            lines.append(f"- {warning}")
        lines.append("")

    lines.append("## Extracted Draft")
    lines.append("")

    for page in result.pages:
        lines.append(f"<!-- PDF_PAGE_BREAK: {page.page_number} -->")
        lines.append("")
        lines.append(f"### Page {page.page_number}")
        lines.append("")
        lines.append(page_text_to_editable_markdown(page))
        lines.append("")

    return "\n".join(lines).rstrip() + "\n"


def write_pdf_text_markdown(
    input_pdf: str | Path,
    output_dir: str | Path,
) -> PdfTextExtractionResult:
    return write_pdf_text_exports(
        input_pdf,
        output_dir,
        export_markdown=True,
        export_docx=False,
    )



def write_pdf_text_docx(
    result: PdfTextExtractionResult,
    output_dir: str | Path,
) -> Path:
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    docx_path = output_dir / f"{safe_pdf_text_stem(result.input_pdf)} - Extracted Text.docx"

    document = Document()
    document.add_heading(f"Editable draft extracted from: {result.input_pdf.stem}", level=1)

    warning = document.add_paragraph()
    warning.add_run("Experimental PDF text recovery. ").bold = True
    warning.add_run(
        "This DOCX is an editable draft, not a polished book conversion. "
        "OCR, layout-perfect reconstruction, footnotes, tables, images, columns, and advanced reflow are not included."
    )

    document.add_paragraph(f"Source PDF: {result.input_pdf}")
    document.add_paragraph(f"Pages: {result.page_count}")
    document.add_paragraph(f"Extracted characters: {result.total_chars}")

    if result.quality_report:
        report = result.quality_report
        document.add_heading("Extraction Quality Report", level=2)
        document.add_paragraph(f"Extraction confidence: {report.confidence}")
        document.add_paragraph(f"OCR needed: {report.ocr_needed}")
        document.add_paragraph(f"Average characters per page: {report.average_chars_per_page:.1f}")
        document.add_paragraph(f"Empty pages: {len(report.empty_pages)}")
        document.add_paragraph(f"Near-empty pages: {len(report.near_empty_pages)}")
        document.add_paragraph(f"Likely scanned/image-only pages: {len(report.likely_scanned_pages)}")
        document.add_paragraph(f"Garbled-text pages: {len(report.garbled_pages)}")

        if report.repeated_line_candidates:
            document.add_paragraph("Repeated line candidates:")
            for candidate in report.repeated_line_candidates[:10]:
                document.add_paragraph(candidate, style="List Bullet")

    if result.warnings:
        document.add_heading("Warnings", level=2)
        for warning_text in result.warnings:
            document.add_paragraph(warning_text, style="List Bullet")

    document.add_heading("Extracted Draft", level=2)

    for page in result.pages:
        document.add_paragraph(f"PDF_PAGE_BREAK: {page.page_number}")
        document.add_heading(f"Page {page.page_number}", level=3)

        editable = page_text_to_editable_markdown(page)
        for block in editable.split("\n\n"):
            value = block.strip()
            if not value:
                continue

            if value.startswith("## "):
                document.add_heading(value[3:].strip(), level=3)
            else:
                document.add_paragraph(value)

    document.save(docx_path)
    result.docx_path = docx_path
    return docx_path


def write_pdf_text_exports(
    input_pdf: str | Path,
    output_dir: str | Path,
    *,
    export_markdown: bool = True,
    export_docx: bool = False,
) -> PdfTextExtractionResult:
    result = extract_pdf_text_pages(input_pdf)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    if export_markdown:
        markdown_path = output_dir / f"{safe_pdf_text_stem(result.input_pdf)} - Extracted Text.md"
        markdown_path.write_text(pdf_text_to_markdown(result), encoding="utf-8")
        result.markdown_path = markdown_path

    if export_docx:
        write_pdf_text_docx(result, output_dir)

    return result


def result_to_dict(result: PdfTextExtractionResult) -> dict[str, Any]:
    return {
        "success": result.success,
        "input_pdf": str(result.input_pdf),
        "page_count": result.page_count,
        "total_chars": result.total_chars,
        "markdown_path": str(result.markdown_path) if result.markdown_path else "",
        "docx_path": str(result.docx_path) if result.docx_path else "",
        "warnings": result.warnings,
        "error": result.error,
        "quality_report": (
            {
                "page_count": result.quality_report.page_count,
                "total_chars": result.quality_report.total_chars,
                "average_chars_per_page": result.quality_report.average_chars_per_page,
                "empty_pages": result.quality_report.empty_pages,
                "near_empty_pages": result.quality_report.near_empty_pages,
                "likely_scanned_pages": result.quality_report.likely_scanned_pages,
                "garbled_pages": result.quality_report.garbled_pages,
                "repeated_line_candidates": result.quality_report.repeated_line_candidates,
                "ocr_needed": result.quality_report.ocr_needed,
                "confidence": result.quality_report.confidence,
            }
            if result.quality_report
            else None
        ),
        "pages": [
            {
                "page_number": page.page_number,
                "char_count": page.char_count,
                "text_preview": page.text[:240],
            }
            for page in result.pages
        ],
    }

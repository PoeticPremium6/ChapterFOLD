from __future__ import annotations

from dataclasses import dataclass, asdict
import json
import math
import re
from pathlib import Path

try:
    from docx import Document
except Exception:  # pragma: no cover
    Document = None


@dataclass(frozen=True)
class SignatureGroup:
    number: int
    start_page: int
    end_page: int
    blank_pages: int = 0

    @property
    def label(self) -> str:
        return f"Signature {self.number}: pages {self.start_page}-{self.end_page}"


def calculate_signature_groups(total_pages: int, pages_per_signature: int) -> list[SignatureGroup]:
    if total_pages <= 0:
        return []
    if pages_per_signature <= 0 or pages_per_signature % 4 != 0:
        raise ValueError("pages_per_signature must be a positive multiple of 4")

    padded_total = int(math.ceil(total_pages / pages_per_signature) * pages_per_signature)
    groups: list[SignatureGroup] = []

    for i, start in enumerate(range(1, padded_total + 1, pages_per_signature), start=1):
        end = min(start + pages_per_signature - 1, padded_total)
        real_end = min(end, total_pages)
        blanks = max(0, end - real_end)
        groups.append(SignatureGroup(i, start, end, blanks))

    return groups


def write_signature_plan_json(
    *,
    output_path: Path,
    total_pages: int,
    pages_per_signature: int,
    groups: list[SignatureGroup],
) -> Path:
    payload = {
        "total_pages": total_pages,
        "pages_per_signature": pages_per_signature,
        "total_signatures": len(groups),
        "blank_pages_added": sum(group.blank_pages for group in groups),
        "signatures": [asdict(group) for group in groups],
    }
    output_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return output_path


def write_signature_plan_markdown(
    *,
    output_path: Path,
    title: str,
    total_pages: int,
    pages_per_signature: int,
    groups: list[SignatureGroup],
) -> Path:
    lines = [
        f"# Signature Plan — {title or 'Book'}",
        "",
        f"- Total pages: {total_pages}",
        f"- Pages per signature: {pages_per_signature}",
        f"- Total signatures: {len(groups)}",
        f"- Blank pages added: {sum(group.blank_pages for group in groups)}",
        "",
        "## Signatures",
        "",
    ]

    for group in groups:
        suffix = f" — {group.blank_pages} blank page(s)" if group.blank_pages else ""
        lines.append(f"- {group.label}{suffix}")

    output_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return output_path


def split_markdown_into_signature_batches(markdown_text: str, groups: list[SignatureGroup]) -> str:
    """Create an editable Markdown companion grouped by signature.

    This is intentionally approximate: Markdown has no final page model, so this
    creates clear signature sections rather than true page-accurate imposition.
    """
    text = markdown_text.strip()
    if not text or not groups:
        return markdown_text

    paragraphs = re.split(r"\n\s*\n", text)
    per_group = max(1, math.ceil(len(paragraphs) / len(groups)))

    lines: list[str] = ["# Signature Batches", ""]
    cursor = 0

    for group in groups:
        lines.extend([f"## {group.label}", ""])
        chunk = paragraphs[cursor : cursor + per_group]
        cursor += per_group
        if chunk:
            lines.append("\n\n".join(chunk).strip())
            lines.append("")
        if group.blank_pages:
            lines.append(f"<!-- {group.blank_pages} blank page(s) added in imposed PDF -->")
            lines.append("")

    if cursor < len(paragraphs):
        lines.extend(["## Overflow", "", "\n\n".join(paragraphs[cursor:]).strip(), ""])

    return "\n".join(lines).strip() + "\n"


def write_signature_markdown(
    *,
    source_markdown: Path,
    output_path: Path,
    groups: list[SignatureGroup],
) -> Path:
    markdown_text = source_markdown.read_text(encoding="utf-8")
    output_path.write_text(split_markdown_into_signature_batches(markdown_text, groups), encoding="utf-8")
    return output_path


def write_signature_docx(
    *,
    output_path: Path,
    title: str,
    groups: list[SignatureGroup],
) -> Path:
    if Document is None:
        raise RuntimeError("python-docx is required for signature DOCX export")

    doc = Document()
    doc.add_heading(f"Signature Batches — {title or 'Book'}", 0)
    doc.add_paragraph(
        "DOCX cannot be physically imposed like a PDF. This companion file lists "
        "the same signature ranges so the editable workflow can track print batches."
    )

    for group in groups:
        doc.add_heading(group.label, level=1)
        if group.blank_pages:
            doc.add_paragraph(f"Blank pages added in imposed PDF: {group.blank_pages}")
        else:
            doc.add_paragraph("No blank pages added in this signature.")

    doc.save(str(output_path))
    return output_path
